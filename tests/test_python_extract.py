import pytest
import os
import json
import fitz
import docx
from python.extract_pdf import extract_pdf
from python.extract_docx import extract_docx

def test_extract_pdf(tmp_path):
    pdf_path = os.path.join(tmp_path, "sample.pdf")
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Hello Legal World! This is a test contract.")
    doc.save(pdf_path)
    doc.close()

    result = extract_pdf(pdf_path)
    assert "error" not in result
    assert result["page_count"] == 1
    assert result["is_scanned"] is False
    assert "Hello Legal World!" in result["full_text"]
    assert len(result["pages"][0]["words"]) > 0

def test_extract_scanned_pdf(tmp_path):
    pdf_path = os.path.join(tmp_path, "scanned.pdf")
    doc = fitz.open()
    doc.new_page()  # empty page
    doc.save(pdf_path)
    doc.close()

    result = extract_pdf(pdf_path)
    assert result["is_scanned"] is True

def test_extract_docx(tmp_path):
    docx_path = os.path.join(tmp_path, "sample.docx")
    doc = docx.Document()
    doc.add_heading("Section 1. Agreement", level=1)
    doc.add_paragraph("The Parties agree to the terms herein.")
    doc.save(docx_path)

    result = extract_docx(docx_path)
    assert "error" not in result
    assert "Agreement" in result["full_text"]
    assert "Parties agree" in result["full_text"]
