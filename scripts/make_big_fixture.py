import os
import fitz  # PyMuPDF
import docx

def create_150_page_pdf(output_path):
    doc = fitz.open()
    for i in range(1, 151):
        page = doc.new_page()
        text = f"MASTER SERVICES AGREEMENT - PAGE {i}\n\n"
        text += f"Section {i}. Obligations and Terms for Clause {i}\n\n"
        text += f"1.{i} The Service Provider agrees to deliver services in accordance with Schedule {i}.\n"
        text += f"The total aggregate liability under Section {i} shall not exceed AED {100000 + i * 1000}.\n"
        text += "All notices and communications shall be provided in writing to the designated address."
        page.insert_text((50, 72), text)
    
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    doc.close()
    print(f"Generated 150-page PDF at: {output_path}")

def create_150_page_docx(output_path):
    doc = docx.Document()
    for i in range(1, 151):
        doc.add_heading(f"Section {i}. Master Services Provision {i}", level=1)
        doc.add_paragraph(
            f"1.{i} The contracting parties hereby agree to terms under Clause {i}. "
            f"The payment obligation for this section shall be AED {50000 + i * 500} net 30 days."
        )
        if i < 150:
            doc.add_page_break()

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Generated 150-page DOCX at: {output_path}")

if __name__ == "__main__":
    fixtures_dir = os.path.join(os.path.dirname(__file__), "..", "fixtures")
    create_150_page_pdf(os.path.join(fixtures_dir, "big_contract.pdf"))
    create_150_page_docx(os.path.join(fixtures_dir, "big_contract.docx"))
