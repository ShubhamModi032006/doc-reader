import sys
import json
import docx

def extract_docx(file_path):
    try:
        doc = docx.Document(file_path)
    except Exception as e:
        return {"error": f"Failed to open DOCX: {str(e)}"}

    full_text_parts = []
    structure = []

    # Extract paragraphs and headings
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        
        style_name = p.style.name if p.style else ""
        is_heading = style_name.startswith("Heading") or text.isupper() and len(text) < 100

        structure.append({
            "type": "heading" if is_heading else "paragraph",
            "text": text,
            "style": style_name
        })
        full_text_parts.append(text)

    # Extract tables
    for t_idx, table in enumerate(doc.tables):
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            table_rows.append(row_cells)
            full_text_parts.append(" | ".join(row_cells))
        
        structure.append({
            "type": "table",
            "table_index": t_idx,
            "rows": table_rows
        })

    full_text = "\n\n".join(full_text_parts)
    
    # Estimate page count roughly (~3000 chars per page, minimum 1)
    estimated_pages = max(1, (len(full_text) + 2999) // 3000)

    return {
        "page_count": estimated_pages,
        "total_chars": len(full_text),
        "is_scanned": False,
        "full_text": full_text,
        "structure": structure
    }

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(json.dumps({"error": "No file path provided"}))
        sys.exit(1)

    docx_path = sys.argv[1]
    result = extract_docx(docx_path)
    print(json.dumps(result, ensure_ascii=False))
